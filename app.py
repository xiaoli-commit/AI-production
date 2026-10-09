from __future__ import annotations

import re
import json
import hashlib
import sqlite3
import threading
import uuid
from datetime import date, datetime, timezone
from difflib import SequenceMatcher
from io import BytesIO
from pathlib import Path
from typing import Any
from urllib.parse import quote

import numpy as np
import pymupdf
from docx import Document
from fastapi import FastAPI, File, Form, HTTPException, Request, UploadFile
from fastapi.responses import FileResponse, StreamingResponse
from openpyxl import load_workbook
from rapidocr_onnxruntime import RapidOCR
from starlette.concurrency import run_in_threadpool


BASE_DIR = Path(__file__).resolve().parent
PAGE_PATH = BASE_DIR / "采购审查平台.html"
DATABASE_PATH = BASE_DIR / "采购审查助手数据.sqlite3"
SAVED_DOCUMENT_DIR = BASE_DIR / "采购审查助手保存文档"
MAX_FILE_BYTES = 20 * 1024 * 1024
MAX_TEXT_CHARS = 300_000
MAX_OCR_PAGES = 100
ALLOWED_EXTENSIONS = {".pdf", ".docx", ".xlsx"}
_ocr_engine: RapidOCR | None = None
_ocr_engine_lock = threading.Lock()
DATE_PATTERNS = (
    re.compile(
        r"(?P<year>19\d{2}|20\d{2})\s*年\s*(?P<month>\d{1,2})\s*月\s*(?P<day>\d{1,2})\s*日?"
    ),
    re.compile(
        r"(?P<year>19\d{2}|20\d{2})(?P<separator>[-/.])"
        r"(?P<month>\d{1,2})(?P=separator)(?P<day>\d{1,2})"
    ),
)

app = FastAPI(title="采购审查助手本地提取服务", docs_url=None, redoc_url=None)

FIELD_RULES = {
    "supplier_qualification": ("供应商资格", r"供应商资格|资格条件|资格要求|资格审查|资质要求|营业执照|信用记录"),
    "technical_requirements": ("技术参数", r"技术参数|技术要求|技术规范|规格型号|性能指标|功能需求|质量标准"),
    "evaluation_method": ("评审办法", r"评审办法|评标办法|评审标准|评分标准|综合评分|价格分|评审因素"),
    "payment_terms": ("付款条件", r"付款|支付|结算|预付款|尾款|付款方式"),
    "delivery_terms": ("交付期限", r"交付期限|交货期限|交货时间|履行期限|服务期限|交付时间|工期"),
    "liability_terms": ("违约责任", r"违约责任|违约金|逾期交付|赔偿损失|违约处理"),
}
PROCUREMENT_METHODS = (
    "公开招标", "邀请招标", "竞争性磋商", "竞争性谈判", "单一来源", "询价采购", "询价",
    "询比采购", "比选采购", "比选", "竞价采购", "竞争性采购",
)
BUDGET_PATTERN = re.compile(
    r"(?:项目预算|采购预算|预算金额|预算总额|预算|最高投标限价|最高限价|控制价)"
    r"[^\d\n]{0,24}(?:人民币|RMB|￥|¥)?\s*([\d,]+(?:\.\d+)?)\s*(亿元|万元|万|元)?"
)
FILENAME_AMOUNT_PATTERN = re.compile(r"(?<![\d.])([\d,]+(?:\.\d+)?)\s*(亿元|万元|万|元)(?![年月日])")


def _connect_database() -> sqlite3.Connection:
    connection = sqlite3.connect(DATABASE_PATH)
    connection.row_factory = sqlite3.Row
    connection.execute(
        "CREATE TABLE IF NOT EXISTS knowledge ("
        "id TEXT PRIMARY KEY, title TEXT NOT NULL, category TEXT NOT NULL, issuer TEXT NOT NULL, "
        "source TEXT NOT NULL, effective_date TEXT NOT NULL, applicability TEXT NOT NULL, "
        "text TEXT NOT NULL, imported_at TEXT NOT NULL)"
    )
    connection.execute(
        "CREATE TABLE IF NOT EXISTS projects ("
        "id TEXT PRIMARY KEY, name TEXT NOT NULL, created_at TEXT NOT NULL, budget REAL, "
        "fields_json TEXT NOT NULL, documents_json TEXT NOT NULL, findings_json TEXT NOT NULL, "
        "comparison_id TEXT NOT NULL DEFAULT '', reviewed INTEGER NOT NULL DEFAULT 0)"
    )
    connection.execute(
        "CREATE TABLE IF NOT EXISTS saved_documents ("
        "id TEXT PRIMARY KEY, client_key TEXT NOT NULL UNIQUE, filename TEXT NOT NULL, file_path TEXT NOT NULL, "
        "size INTEGER NOT NULL, sha256 TEXT NOT NULL, uploaded_at TEXT NOT NULL, extracted_text TEXT NOT NULL DEFAULT '', "
        "extraction_json TEXT NOT NULL DEFAULT '{}', extracted_at TEXT NOT NULL DEFAULT '')"
    )
    connection.commit()
    return connection


def _now() -> str:
    return datetime.now(timezone.utc).astimezone().isoformat(timespec="seconds")


def _read_json(value: str) -> Any:
    return json.loads(value)


def _document_categories(filename: str, text: str) -> list[str]:
    sample = f"{filename} {text}"
    patterns = (
        ("评审报告", r"评审报告|评标报告|评分汇总|评审委员会"),
        ("采购合同", r"采购合同|合同条款|合同文本|合同签订|本合同|甲方|乙方"),
        ("技术规范", r"技术规范|技术参数|技术要求|技术部分|性能指标|规格型号"),
        ("招投标文件", r"招标文件|投标文件|招标公告|投标人|采购文件|比选文件|询比文件|遴选文件|磋商文件"),
    )
    categories = [category for category, pattern in patterns if re.search(pattern, sample)]
    return categories or ["采购文档"]


def _money_to_yuan(amount: str, unit: str | None) -> float:
    multiplier = {"亿元": 100_000_000, "万元": 10_000, "万": 10_000, "元": 1}.get(unit or "元", 1)
    return float(amount.replace(",", "")) * multiplier


def _extract_key_fields(documents: list[dict[str, Any]], budget_override: float | None) -> dict[str, Any]:
    fields: dict[str, Any] = {
        "procurement_method": {"label": "采购方式", "values": []},
        "budget": {"label": "项目预算", "values": []},
        **{key: {"label": label, "values": []} for key, (label, _) in FIELD_RULES.items()},
    }
    all_text = "\n".join(document.get("text", "") for document in documents if not document.get("error"))
    for document in documents:
        if document.get("error"):
            continue
        filename = document["filename"]
        blocks = [{"source": "文件名", "text": filename}]
        blocks.extend(document.get("blocks") or [{"source": "全文", "text": document.get("text", "")}])
        for block in blocks:
            text = block.get("text", "")
            for method in PROCUREMENT_METHODS:
                if method in text and not any(item["value"] == method for item in fields["procurement_method"]["values"]):
                    fields["procurement_method"]["values"].append(
                        {"value": method, "source": filename, "location": block.get("source", "全文"), "snippet": _context(text, text.find(method), text.find(method) + len(method), 90)}
                    )
            for match in BUDGET_PATTERN.finditer(text):
                amount = _money_to_yuan(match.group(1), match.group(2))
                fields["budget"]["values"].append(
                    {"value": amount, "display": match.group(0).strip(), "source": filename, "location": block.get("source", "全文"), "snippet": _context(text, match.start(), match.end(), 100)}
                )
            if block.get("source") == "文件名" and not list(BUDGET_PATTERN.finditer(text)):
                for match in FILENAME_AMOUNT_PATTERN.finditer(text):
                    amount = _money_to_yuan(match.group(1), match.group(2))
                    fields["budget"]["values"].append(
                        {"value": amount, "display": match.group(0).strip(), "source": filename, "location": "文件名", "snippet": text}
                    )
            for key, (label, pattern) in FIELD_RULES.items():
                for match in re.finditer(pattern, text):
                    snippet = _context(text, match.start(), match.end(), 180)
                    value = {"value": snippet, "source": filename, "location": block.get("source", "全文"), "snippet": snippet}
                    if not any(item["source"] == filename and item["snippet"] == snippet for item in fields[key]["values"]):
                        fields[key]["values"].append(value)
                    if len(fields[key]["values"]) >= 20:
                        break

    if budget_override is not None:
        fields["budget"]["manual_value"] = budget_override
    fields["document_types"] = list(dict.fromkeys(
        category
        for document in documents if not document.get("error")
        for category in _document_categories(document["filename"], document.get("text", ""))
    ))
    return fields


def _issue(
    severity: str,
    category: str,
    title: str,
    message: str,
    source: str = "",
    location: str = "",
    context: str = "",
    suggestion: str = "",
) -> dict[str, Any]:
    return {
        "id": uuid.uuid4().hex,
        "severity": severity,
        "category": category,
        "title": title,
        "message": message,
        "source": source,
        "location": location,
        "context": context,
        "suggestion": suggestion,
        "citations": [],
        "review_status": "待复核",
    }


def _knowledge_citations(category: str, knowledge: list[dict[str, Any]]) -> list[dict[str, str]]:
    topic_terms = {
        "日期": r"日期|期限|履行|交付|有效期",
        "预算": r"预算|金额|限价|采购计划",
        "采购方式": r"招标|采购方式|磋商|谈判|询价|单一来源",
        "供应商资格": r"供应商|资格|资质|信用|条件",
        "评审办法": r"评审|评标|评分|评委|评标委员会",
        "技术参数": r"技术|参数|规格|性能|质量",
        "付款条件": r"付款|支付|结算|预付款",
        "交付期限": r"交付|交货|服务期限|履行期限|工期",
        "违约责任": r"违约|责任|赔偿|损失",
        "文档完整性": r"采购|合同|招标|投标|评审",
        "历史差异": r"采购|合同|招标|投标|评审",
    }.get(category, r"采购|合同|招标")
    citations = []
    for item in knowledge:
        match = re.search(topic_terms, item["text"])
        if not match:
            continue
        citations.append(
            {
                "title": item["title"],
                "issuer": item["issuer"],
                "source": item["source"],
                "effective_date": item["effective_date"],
                "applicability": item["applicability"],
                "excerpt": _context(item["text"], match.start(), match.end(), 90),
            }
        )
        if len(citations) == 3:
            break
    return citations


def _review_documents(
    documents: list[dict[str, Any]],
    fields: dict[str, Any],
    knowledge: list[dict[str, Any]],
    budget_override: float | None,
    previous: dict[str, Any] | None,
) -> list[dict[str, Any]]:
    findings: list[dict[str, Any]] = []
    valid_documents = [document for document in documents if not document.get("error")]
    categories = set(fields.get("document_types", []))

    for document in documents:
        for item in document.get("issues", []):
            severity = "高" if item["type"] in {"invalid_date", "date_order"} else "中"
            findings.append(
                _issue(
                    severity, "日期", item["message"], item["message"], document["filename"],
                    item.get("source", ""), item.get("context", ""), "核对原文日期及相关期限，并统一日期格式。",
                )
            )

    contract = "采购合同" in categories
    tender = "招投标文件" in categories
    required = []
    if tender:
        required.extend(["procurement_method", "supplier_qualification", "evaluation_method", "technical_requirements"])
    if contract:
        required.extend(["payment_terms", "delivery_terms", "liability_terms"])
    for key in dict.fromkeys(required):
        if not fields[key]["values"]:
            label = fields[key]["label"]
            severity = "高" if contract and key in {"payment_terms", "delivery_terms", "liability_terms"} else "中"
            findings.append(
                _issue(
                    severity, label, f"未识别到{label}",
                    f"在当前提取文本中未识别到明确的{label}表述；也可能存在不同说法或文本提取遗漏。",
                    suggestion=f"人工检查原文是否包含{label}，必要时补充并明确具体内容。",
                )
            )
        for value in fields[key]["values"]:
            snippet = value.get("snippet", value.get("value", ""))
            if re.search(r"(?:待填|待定|待补充|未填写|暂无|TBD|\[\s*\]|_{2,}|X{2,})", snippet, re.IGNORECASE):
                findings.append(
                    _issue(
                        "高" if contract and key in {"payment_terms", "delivery_terms", "liability_terms"} else "中",
                        fields[key]["label"], f"{fields[key]['label']}可能仍有未填写占位内容",
                        f"识别到{fields[key]['label']}附近含有空白或待填写占位符，请核对是否为未完成条款。",
                        value.get("source", ""), value.get("location", ""), snippet,
                        f"补充完整的{fields[key]['label']}内容，并复核与采购文件及合同其他条款的一致性。",
                    )
                )

    budgets = fields["budget"]["values"]
    if len(budgets) > 1:
        amounts = [item["value"] for item in budgets]
        low, high = min(amounts), max(amounts)
        if low > 0 and (high - low) / low > 0.01:
            findings.append(
                _issue(
                    "高", "预算", "不同文件中的预算金额不一致",
                    f"检测到预算金额从 {low:,.2f} 元到 {high:,.2f} 元，偏差超过 1%。请确认采用哪个金额。",
                    budgets[-1]["source"], budgets[-1]["location"], budgets[-1]["snippet"], "核对采购预算批复、公告及合同金额口径。",
                )
            )
    if budget_override is not None:
        detected = [item["value"] for item in budgets]
        if detected and min(abs(value - budget_override) for value in detected) / max(budget_override, 1) > 0.01:
            findings.append(
                _issue(
                    "高", "预算", "录入预算与文档金额不一致",
                    f"录入预算为 {budget_override:,.2f} 元，文档识别金额为 {min(detected, key=lambda value: abs(value-budget_override)):,.2f} 元。",
                    suggestion="核对预算输入及原文件中的预算、限价口径。",
                )
            )

    if previous:
        old_fields = previous["fields"]
        old_budget = old_fields.get("budget", {}).get("manual_value")
        if old_budget is None:
            old_values = old_fields.get("budget", {}).get("values", [])
            old_budget = old_values[0]["value"] if old_values else None
        new_budget = budget_override
        if new_budget is None and budgets:
            new_budget = budgets[0]["value"]
        if old_budget and new_budget and abs(new_budget-old_budget) / old_budget > 0.3:
            findings.append(
                _issue(
                    "中", "历史差异", "预算与选定历史项目差异较大",
                    f"当前预算 {new_budget:,.2f} 元，历史项目“{previous['name']}”为 {old_budget:,.2f} 元，变化超过 30%。",
                    suggestion="核实项目范围、采购数量和预算依据是否可比。",
                )
            )
        for key in ("procurement_method", "payment_terms", "delivery_terms", "liability_terms", "technical_requirements"):
            current_values = fields.get(key, {}).get("values", [])
            prior_values = old_fields.get(key, {}).get("values", [])
            if not current_values or not prior_values:
                continue
            current_text = " ".join(item["value"] for item in current_values)[:4000]
            prior_text = " ".join(item["value"] for item in prior_values)[:4000]
            if current_text != prior_text and SequenceMatcher(None, current_text, prior_text).ratio() < 0.55:
                label = fields[key]["label"]
                findings.append(
                    _issue(
                        "中", "历史差异", f"{label}与选定历史项目差异明显",
                        f"当前文档中的{label}与历史项目“{previous['name']}”文本差异较大；差异本身不代表违规。",
                        suggestion="确认两个项目范围可比后，复核变化原因及审批依据。",
                    )
                )

    for finding in findings:
        finding["citations"] = _knowledge_citations(finding["category"], knowledge)
    if not knowledge and findings:
        findings.append(
            _issue(
                "提示", "知识库", "未配置制度来源",
                "本机制度知识库尚无已导入且可匹配的制度文件；当前提示只依据文本规则生成，没有法规条款引用。",
                suggestion="在“制度知识库”导入经核验的法规或所内制度后重新审查。",
            )
        )
    severity_order = {"高": 0, "中": 1, "低": 2, "提示": 3}
    findings.sort(key=lambda item: severity_order.get(item["severity"], 4))
    return findings


@app.get("/")
def index() -> FileResponse:
    return FileResponse(PAGE_PATH)


@app.get("/api/health")
def health() -> dict[str, Any]:
    connection = _connect_database()
    knowledge_count = connection.execute("SELECT COUNT(*) FROM knowledge").fetchone()[0]
    connection.close()
    return {
        "ok": True,
        "max_file_bytes": MAX_FILE_BYTES,
        "formats": sorted(ALLOWED_EXTENSIONS),
        "knowledge_count": knowledge_count,
        "ollama_configured": False,
    }


def _date_matches(text: str) -> list[dict[str, Any]]:
    matches: list[dict[str, Any]] = []
    occupied: list[tuple[int, int]] = []
    for pattern in DATE_PATTERNS:
        for match in pattern.finditer(text):
            if any(match.start() < end and match.end() > start for start, end in occupied):
                continue
            occupied.append((match.start(), match.end()))
            year, month, day = (int(match.group(key)) for key in ("year", "month", "day"))
            raw = match.group(0)
            separator = match.groupdict().get("separator")
            try:
                parsed = date(year, month, day)
                normalized = parsed.isoformat()
                valid = True
            except ValueError:
                normalized = None
                valid = False
            matches.append(
                {
                    "start": match.start(),
                    "end": match.end(),
                    "raw": raw,
                    "normalized": normalized,
                    "valid": valid,
                    "separator": separator,
                    "month": month,
                    "day": day,
                }
            )
    return sorted(matches, key=lambda item: item["start"])


def _context(text: str, start: int, end: int, width: int = 45) -> str:
    left = max(0, start - width)
    right = min(len(text), end + width)
    return text[left:right].replace("\n", " ").strip()


def analyze_dates(blocks: list[dict[str, str]]) -> tuple[list[dict[str, Any]], list[dict[str, str]]]:
    dates: list[dict[str, Any]] = []
    issues: list[dict[str, str]] = []
    seen_issues: set[tuple[str, str, str]] = set()

    def add_issue(kind: str, message: str, source: str, context: str) -> None:
        key = (kind, source, context)
        if key not in seen_issues:
            seen_issues.add(key)
            issues.append({"type": kind, "message": message, "source": source, "context": context})

    for block in blocks:
        text = block["text"]
        source = block["source"]
        matches = _date_matches(text)
        for match in matches:
            context = _context(text, match["start"], match["end"])
            dates.append(
                {
                    "text": match["raw"],
                    "normalized": match["normalized"],
                    "valid": match["valid"],
                    "source": source,
                    "context": context,
                }
            )
            if not match["valid"]:
                add_issue("invalid_date", f"日期 {match['raw']} 不是有效日历日期。", source, context)
            elif match["separator"] in {"/", "."}:
                ambiguous = match["separator"] == "/" and match["month"] <= 12 and match["day"] <= 12
                add_issue(
                    "ambiguous_date" if ambiguous else "date_format",
                    f"日期 {match['raw']} 使用{'易混淆的' if ambiguous else '非推荐的'}分隔格式，建议统一为 YYYY年MM月DD日。",
                    source,
                    context,
                )

        for clause in re.split(r"[。；;\n]+", text):
            if not re.search(r"期限|有效期|服务期|履行期|自.{0,12}(?:至|到)|从.{0,12}(?:至|到)", clause):
                continue
            clause_dates = [item for item in _date_matches(clause) if item["valid"]]
            if len(clause_dates) >= 2 and clause_dates[0]["normalized"] > clause_dates[1]["normalized"]:
                add_issue(
                    "date_order",
                    "期限条款中的起止日期先后倒置，请核对。",
                    source,
                    clause.strip()[:180],
                )

        for placeholder in re.finditer(
            r"(?:签订日期|签署日期|生效日期|截止日期|日期|期限|交付时间|交货时间)[^。；;\n]{0,20}(?:待填|待定|填写|\[\s*\]|_{2,}|X{2,})",
            text,
            flags=re.IGNORECASE,
        ):
            add_issue(
                "date_placeholder",
                "日期字段似乎仍含待填写占位内容。",
                source,
                _context(text, placeholder.start(), placeholder.end()),
            )

    if not dates and any(re.search(r"合同|协议|签订|日期|期限", block["text"]) for block in blocks):
        add_issue("date_missing", "文本包含合同或日期相关内容，但未识别到标准数字日期，请人工核对。", "全文", "")
    return dates[:500], issues[:300]


def _append_block(blocks: list[dict[str, str]], source: str, text: str) -> int:
    remaining = MAX_TEXT_CHARS - sum(len(block["text"]) for block in blocks)
    if remaining <= 0:
        return 0
    content = text.strip()[:remaining]
    if content:
        blocks.append({"source": source, "text": content})
    return len(content)


def _get_ocr_engine() -> RapidOCR:
    global _ocr_engine
    if _ocr_engine is None:
        with _ocr_engine_lock:
            if _ocr_engine is None:
                _ocr_engine = RapidOCR()
    return _ocr_engine


def _extract_pdf(raw: bytes, enable_ocr: bool = True) -> tuple[list[dict[str, str]], int, int, bool]:
    document = pymupdf.open(stream=raw, filetype="pdf")
    page_count = len(document)
    blocks: list[dict[str, str]] = []
    ocr_pages = 0
    ocr_limit_reached = False
    try:
        for page_number, page in enumerate(document, start=1):
            text = page.get_text("text") or ""
            if not text.strip():
                if not enable_ocr:
                    continue
                if ocr_pages >= MAX_OCR_PAGES:
                    ocr_limit_reached = True
                    continue
                pixmap = page.get_pixmap(
                    matrix=pymupdf.Matrix(1.5, 1.5),
                    colorspace=pymupdf.csRGB,
                    alpha=False,
                )
                image = np.frombuffer(pixmap.samples, dtype=np.uint8).reshape(
                    pixmap.height, pixmap.width, pixmap.n
                )
                result, _ = _get_ocr_engine()(image)
                text = " ".join(item[1] for item in (result or []) if len(item) > 1)
                ocr_pages += 1
            if text.strip():
                _append_block(blocks, f"第 {page_number} 页", text)
    finally:
        document.close()
    return blocks, page_count, ocr_pages, ocr_limit_reached


def _extract_docx(raw: bytes) -> tuple[list[dict[str, str]], int]:
    document = Document(BytesIO(raw))
    blocks: list[dict[str, str]] = []
    for index, paragraph in enumerate(document.paragraphs, start=1):
        if paragraph.text.strip():
            _append_block(blocks, f"段落 {index}", paragraph.text)
    for table_index, table in enumerate(document.tables, start=1):
        for row_index, row in enumerate(table.rows, start=1):
            values = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            line = " | ".join(value for value in values if value)
            if line:
                _append_block(blocks, f"表格 {table_index} 行 {row_index}", line)
    return blocks, len(blocks)


def _extract_xlsx(raw: bytes) -> tuple[list[dict[str, str]], int]:
    workbook = load_workbook(BytesIO(raw), read_only=True, data_only=True)
    blocks: list[dict[str, str]] = []
    for worksheet in workbook.worksheets:
        for row_index, row in enumerate(worksheet.iter_rows(), start=1):
            values: list[str] = []
            for cell in row:
                value = cell.value
                if isinstance(value, (datetime, date)):
                    values.append(value.strftime("%Y-%m-%d"))
                elif value is not None:
                    values.append(str(value).strip())
            line = " | ".join(value for value in values if value)
            if line:
                _append_block(blocks, f"工作表 {worksheet.title} 第 {row_index} 行", line)
    workbook.close()
    return blocks, len(blocks)


def extract_document(filename: str, raw: bytes, enable_ocr: bool = True) -> dict[str, Any]:
    extension = Path(filename).suffix.lower()
    if extension not in ALLOWED_EXTENSIONS:
        message = "暂不支持此格式；旧版 .doc/.xls 请另存为 .docx/.xlsx。"
        return {"filename": filename, "error": message}
    if len(raw) > MAX_FILE_BYTES:
        return {"filename": filename, "error": "文件超过 20MB 限制。"}

    try:
        ocr_pages = 0
        ocr_limit_reached = False
        if extension == ".pdf":
            blocks, source_count, ocr_pages, ocr_limit_reached = _extract_pdf(raw, enable_ocr)
        elif extension == ".docx":
            blocks, source_count = _extract_docx(raw)
        else:
            blocks, source_count = _extract_xlsx(raw)
    except Exception as error:
        return {"filename": filename, "error": f"文档解析失败：{type(error).__name__}。请确认文件未损坏且格式正确。"}

    text = "\n\n".join(block["text"] for block in blocks)
    dates, issues = analyze_dates(blocks)
    scanned_pdf = extension == ".pdf" and not text.strip()
    if scanned_pdf:
        issues.append(
            {
                "type": "ocr_required",
                "message": "未识别到可用文本；文件可能是空白页、低清扫描件或图片内容。请人工核对。",
                "source": "全文",
                "context": "",
            }
        )
    if ocr_limit_reached:
        issues.append(
            {
                "type": "ocr_page_limit",
                "message": f"扫描页 OCR 已达到每份文档 {MAX_OCR_PAGES} 页上限，后续扫描页未识别。",
                "source": "全文",
                "context": "",
            }
        )
    return {
        "filename": filename,
        "format": extension[1:].upper(),
        "char_count": len(text),
        "source_count": source_count,
        "ocr_pages": ocr_pages,
        "ocr_limit_reached": ocr_limit_reached,
        "truncated": sum(len(block["text"]) for block in blocks) >= MAX_TEXT_CHARS,
        "text": text,
        "blocks": blocks,
        "dates": dates,
        "issues": issues,
        "error": None,
    }


@app.post("/api/extract")
async def extract(files: list[UploadFile] = File(...)) -> dict[str, Any]:
    results: list[dict[str, Any]] = []
    for upload in files:
        filename = Path(upload.filename or "未命名文件").name
        raw = await upload.read(MAX_FILE_BYTES + 1)
        await upload.close()
        result = await run_in_threadpool(extract_document, filename, raw)
        result.pop("blocks", None)
        results.append(result)
    return {"files": results}


@app.post("/api/prefill")
async def prefill(files: list[UploadFile] = File(...)) -> dict[str, Any]:
    documents: list[dict[str, Any]] = []
    first_filename = ""
    for upload in files:
        filename = Path(upload.filename or "未命名文件").name
        if not first_filename:
            first_filename = filename
        raw = await upload.read(MAX_FILE_BYTES + 1)
        await upload.close()
        documents.append(await run_in_threadpool(extract_document, filename, raw, False))
    if not documents:
        raise HTTPException(status_code=400, detail="请至少选择一个文档。")
    fields = _extract_key_fields(documents, None)
    budget_values = fields["budget"]["values"]
    return {
        "project_name": Path(first_filename).stem,
        "budget": budget_values[0]["value"] if budget_values else None,
        "budget_display": budget_values[0]["display"] if budget_values else "",
        "budget_source": budget_values[0]["source"] if budget_values else "",
        "document_types": fields["document_types"],
    }


def _saved_document_public(row: sqlite3.Row) -> dict[str, Any]:
    return {
        "id": row["id"],
        "client_key": row["client_key"],
        "filename": row["filename"],
        "size": row["size"],
        "sha256": row["sha256"],
        "uploaded_at": row["uploaded_at"],
        "has_extracted_content": bool(row["extracted_at"]),
        "extracted_at": row["extracted_at"],
    }


@app.get("/api/saved-documents")
def list_saved_documents() -> dict[str, Any]:
    connection = _connect_database()
    rows = connection.execute(
        "SELECT id,client_key,filename,size,sha256,uploaded_at,extracted_at "
        "FROM saved_documents ORDER BY uploaded_at DESC"
    ).fetchall()
    connection.close()
    return {"items": [_saved_document_public(row) for row in rows]}


@app.post("/api/saved-documents")
async def save_uploaded_document(
    file: UploadFile = File(...),
    client_key: str = Form(...),
) -> dict[str, Any]:
    filename = Path(file.filename or "未命名文件").name
    extension = Path(filename).suffix.lower()
    if extension not in ALLOWED_EXTENSIONS:
        raise HTTPException(status_code=400, detail="仅支持 PDF、DOCX、XLSX 文件。")
    raw = await file.read(MAX_FILE_BYTES + 1)
    await file.close()
    if len(raw) > MAX_FILE_BYTES:
        raise HTTPException(status_code=413, detail="文件超过 20MB 限制。")
    file_hash = hashlib.sha256(raw).hexdigest()
    client_key = client_key.strip()[:300]
    if not client_key:
        raise HTTPException(status_code=400, detail="缺少文件标识，无法保存。")

    connection = _connect_database()
    existing = connection.execute("SELECT * FROM saved_documents WHERE client_key=?", (client_key,)).fetchone()
    if existing and existing["sha256"] == file_hash:
        result = _saved_document_public(existing)
        connection.close()
        return result

    document_id = existing["id"] if existing else uuid.uuid4().hex
    SAVED_DOCUMENT_DIR.mkdir(parents=True, exist_ok=True)
    destination = SAVED_DOCUMENT_DIR / f"{document_id}{extension}"
    destination.write_bytes(raw)
    uploaded_at = _now()
    if existing:
        old_path = Path(existing["file_path"])
        connection.execute(
            "UPDATE saved_documents SET filename=?,file_path=?,size=?,sha256=?,uploaded_at=?,extracted_text='',extraction_json='{}',extracted_at='' "
            "WHERE id=?",
            (filename, str(destination), len(raw), file_hash, uploaded_at, document_id),
        )
        if old_path != destination and old_path.is_file():
            old_path.unlink()
    else:
        connection.execute(
            "INSERT INTO saved_documents (id,client_key,filename,file_path,size,sha256,uploaded_at) VALUES (?,?,?,?,?,?,?)",
            (document_id, client_key, filename, str(destination), len(raw), file_hash, uploaded_at),
        )
    connection.commit()
    row = connection.execute("SELECT * FROM saved_documents WHERE id=?", (document_id,)).fetchone()
    result = _saved_document_public(row)
    connection.close()
    return result


@app.put("/api/saved-documents/{document_id}/content")
async def save_extracted_document_content(document_id: str, request: Request) -> dict[str, Any]:
    payload = await request.json()
    extracted_text = payload.get("text", "")
    if not isinstance(extracted_text, str):
        raise HTTPException(status_code=400, detail="提取文本格式无效。")
    extracted_text = extracted_text[:MAX_TEXT_CHARS]
    extraction = {
        "format": payload.get("format", ""),
        "dates": payload.get("dates", [])[:500],
        "issues": payload.get("issues", [])[:300],
        "char_count": min(int(payload.get("char_count", len(extracted_text))), MAX_TEXT_CHARS),
        "source_count": int(payload.get("source_count", 0)),
        "ocr_pages": int(payload.get("ocr_pages", 0)),
    }
    connection = _connect_database()
    cursor = connection.execute(
        "UPDATE saved_documents SET extracted_text=?,extraction_json=?,extracted_at=? WHERE id=?",
        (extracted_text, json.dumps(extraction, ensure_ascii=False), _now(), document_id),
    )
    if not cursor.rowcount:
        connection.close()
        raise HTTPException(status_code=404, detail="未找到已保存文件。请先保存原始文档。")
    connection.commit()
    row = connection.execute(
        "SELECT id,client_key,filename,size,sha256,uploaded_at,extracted_at FROM saved_documents WHERE id=?",
        (document_id,),
    ).fetchone()
    connection.close()
    return _saved_document_public(row)


@app.get("/api/saved-documents/{document_id}/file")
def download_saved_document(document_id: str) -> FileResponse:
    connection = _connect_database()
    row = connection.execute("SELECT filename,file_path FROM saved_documents WHERE id=?", (document_id,)).fetchone()
    connection.close()
    if not row or not Path(row["file_path"]).is_file():
        raise HTTPException(status_code=404, detail="未找到已保存文件。")
    return FileResponse(row["file_path"], filename=row["filename"])


def _knowledge_rows() -> list[dict[str, Any]]:
    connection = _connect_database()
    rows = connection.execute(
        "SELECT id, title, category, issuer, source, effective_date, applicability, text, imported_at "
        "FROM knowledge ORDER BY imported_at DESC"
    ).fetchall()
    connection.close()
    return [dict(row) for row in rows]


def _project_row(row: sqlite3.Row) -> dict[str, Any]:
    return {
        "id": row["id"],
        "name": row["name"],
        "created_at": row["created_at"],
        "budget": row["budget"],
        "fields": _read_json(row["fields_json"]),
        "documents": _read_json(row["documents_json"]),
        "findings": _read_json(row["findings_json"]),
        "comparison_id": row["comparison_id"],
        "reviewed": bool(row["reviewed"]),
    }


def _get_project(project_id: str) -> dict[str, Any] | None:
    connection = _connect_database()
    row = connection.execute("SELECT * FROM projects WHERE id=?", (project_id,)).fetchone()
    connection.close()
    return _project_row(row) if row else None


def list_knowledge() -> dict[str, Any]:
    items = _knowledge_rows()
    return {"items": [{key: value for key, value in item.items() if key != "text"} for item in items]}


@app.post("/api/knowledge/import")
async def import_knowledge(
    file: UploadFile = File(...),
    title: str = Form(...),
    category: str = Form(...),
    issuer: str = Form(...),
    source: str = Form(...),
    effective_date: str = Form(""),
    applicability: str = Form(...),
) -> dict[str, Any]:
    filename = Path(file.filename or "制度文件").name
    raw = await file.read(MAX_FILE_BYTES + 1)
    await file.close()
    result = await run_in_threadpool(extract_document, filename, raw)
    if result.get("error"):
        raise HTTPException(status_code=400, detail=result["error"])
    if not result.get("text", "").strip():
        raise HTTPException(status_code=400, detail="未从制度文件中提取到文本，不能加入可引用知识库。")
    if not all(value.strip() for value in (title, category, issuer, source, applicability)):
        raise HTTPException(status_code=400, detail="制度名称、主题、发布方、来源和适用范围均为必填项。")
    item = {
        "id": uuid.uuid4().hex,
        "title": title.strip(),
        "category": category.strip(),
        "issuer": issuer.strip(),
        "source": source.strip(),
        "effective_date": effective_date.strip(),
        "applicability": applicability.strip(),
        "text": result["text"],
        "imported_at": _now(),
    }
    connection = _connect_database()
    connection.execute(
        "INSERT INTO knowledge (id,title,category,issuer,source,effective_date,applicability,text,imported_at) "
        "VALUES (:id,:title,:category,:issuer,:source,:effective_date,:applicability,:text,:imported_at)",
        item,
    )
    connection.commit()
    connection.close()
    return {key: value for key, value in item.items() if key != "text"}


@app.delete("/api/knowledge/{item_id}")
def delete_knowledge(item_id: str) -> dict[str, bool]:
    connection = _connect_database()
    cursor = connection.execute("DELETE FROM knowledge WHERE id=?", (item_id,))
    connection.commit()
    connection.close()
    if not cursor.rowcount:
        raise HTTPException(status_code=404, detail="未找到该制度材料。")
    return {"deleted": True}


@app.post("/api/review")
async def review(
    files: list[UploadFile] = File(...),
    project_name: str = Form(""),
    budget: str = Form(""),
    compare_id: str = Form(""),
) -> dict[str, Any]:
    try:
        budget_override = float(budget.replace(",", "")) if budget.strip() else None
        if budget_override is not None and budget_override <= 0:
            raise ValueError
    except ValueError:
        raise HTTPException(status_code=400, detail="预算需填写大于 0 的金额（元）。") from None

    documents: list[dict[str, Any]] = []
    for upload in files:
        filename = Path(upload.filename or "未命名文件").name
        raw = await upload.read(MAX_FILE_BYTES + 1)
        await upload.close()
        documents.append(await run_in_threadpool(extract_document, filename, raw))

    valid_documents = [item for item in documents if not item.get("error")]
    if not valid_documents:
        raise HTTPException(status_code=400, detail="没有可审查的有效文档。")
    fields = _extract_key_fields(documents, budget_override)
    previous = _get_project(compare_id) if compare_id else None
    if compare_id and previous is None:
        raise HTTPException(status_code=404, detail="未找到指定的历史项目。")
    findings = _review_documents(documents, fields, _knowledge_rows(), budget_override, previous)
    project_id = uuid.uuid4().hex
    created_at = _now()
    resolved_name = project_name.strip() or Path(valid_documents[0]["filename"]).stem
    detected_budgets = fields["budget"]["values"]
    stored_budget = budget_override if budget_override is not None else (detected_budgets[0]["value"] if detected_budgets else None)
    connection = _connect_database()
    connection.execute(
        "INSERT INTO projects (id,name,created_at,budget,fields_json,documents_json,findings_json,comparison_id,reviewed) "
        "VALUES (?,?,?,?,?,?,?,?,0)",
        (
            project_id,
            resolved_name,
            created_at,
            stored_budget,
            json.dumps(fields, ensure_ascii=False),
            json.dumps(documents, ensure_ascii=False),
            json.dumps(findings, ensure_ascii=False),
            compare_id,
        ),
    )
    connection.commit()
    connection.close()
    return {
        "id": project_id,
        "name": resolved_name,
        "created_at": created_at,
        "fields": fields,
        "findings": findings,
        "documents": [
            {
                key: document.get(key)
                for key in ("filename", "format", "char_count", "source_count", "ocr_pages", "truncated", "error", "text", "dates", "issues")
            }
            for document in documents
        ],
        "comparison": {"id": previous["id"], "name": previous["name"]} if previous else None,
    }


@app.get("/api/history")
def history() -> dict[str, Any]:
    connection = _connect_database()
    rows = connection.execute(
        "SELECT id,name,created_at,budget,findings_json,comparison_id,reviewed FROM projects ORDER BY created_at DESC"
    ).fetchall()
    connection.close()
    items = []
    for row in rows:
        findings = _read_json(row["findings_json"])
        counts = {level: sum(item["severity"] == level for item in findings) for level in ("高", "中", "低", "提示")}
        items.append(
            {
                "id": row["id"],
                "name": row["name"],
                "created_at": row["created_at"],
                "budget": row["budget"],
                "finding_count": len(findings),
                "risk_counts": counts,
                "comparison_id": row["comparison_id"],
                "reviewed": bool(row["reviewed"]),
            }
        )
    return {"items": items}


@app.get("/api/history/{project_id}")
def history_item(project_id: str) -> dict[str, Any]:
    project = _get_project(project_id)
    if not project:
        raise HTTPException(status_code=404, detail="未找到该历史项目。")
    project["documents"] = [
        {key: value for key, value in document.items() if key not in {"text", "blocks"}}
        for document in project["documents"]
    ]
    return project


@app.post("/api/history/{project_id}/review")
async def save_review(project_id: str, request: Request) -> dict[str, Any]:
    project = _get_project(project_id)
    if not project:
        raise HTTPException(status_code=404, detail="未找到该审查项目。")
    payload = await request.json()
    findings = payload.get("findings")
    if not isinstance(findings, list) or len(findings) > 500:
        raise HTTPException(status_code=400, detail="审查问题清单格式无效。")
    previous_by_id = {item["id"]: item for item in project["findings"]}
    editable = []
    for item in findings:
        if not isinstance(item, dict) or item.get("id") not in previous_by_id:
            continue
        saved = previous_by_id[item["id"]]
        for key in ("title", "message", "suggestion", "review_status", "severity"):
            if key in item:
                saved[key] = str(item[key])[:3000]
        editable.append(saved)
    connection = _connect_database()
    connection.execute(
        "UPDATE projects SET findings_json=?,reviewed=1 WHERE id=?",
        (json.dumps(editable, ensure_ascii=False), project_id),
    )
    connection.commit()
    connection.close()
    return {"saved": True, "reviewed": True, "findings": editable}


@app.get("/api/history/{project_id}/report.docx")
def report(project_id: str) -> StreamingResponse:
    project = _get_project(project_id)
    if not project:
        raise HTTPException(status_code=404, detail="未找到该审查项目。")
    report_document = Document()
    report_document.add_heading("采购文档审查报告", 0)
    report_document.add_paragraph(f"项目名称：{project['name']}")
    report_document.add_paragraph(f"审查时间：{project['created_at']}")
    report_document.add_paragraph(
        f"预算金额：{project['budget']:,.2f} 元" if project["budget"] is not None else "预算金额：未识别/未录入"
    )
    if project["comparison_id"]:
        prior = _get_project(project["comparison_id"])
        if prior:
            report_document.add_paragraph(f"对比项目：{prior['name']}")
    report_document.add_paragraph("本报告为本地规则辅助结果，需人工复核，不构成法律意见。")
    report_document.add_heading("一、关键项目信息", level=1)
    table = report_document.add_table(rows=1, cols=2)
    table.style = "Light Shading Accent 1"
    table.rows[0].cells[0].text = "信息项"
    table.rows[0].cells[1].text = "提取内容/来源"
    for key, field in project["fields"].items():
        if key == "document_types":
            value = "、".join(field)
            label = "文档类型"
        else:
            label = field["label"]
            values = field.get("values", [])
            value = "；".join(
                f"{item.get('display', item.get('value', ''))}（{item.get('source', '')} {item.get('location', '')}）"
                for item in values[:5]
            ) or "未识别"
            if field.get("manual_value") is not None:
                value = f"人工录入：{field['manual_value']:,.2f} 元；" + value
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = value[:1500]

    report_document.add_heading("二、问题与风险清单", level=1)
    if not project["findings"]:
        report_document.add_paragraph("本次规则检查未生成问题项；仍需人工复核全文。")
    for index, finding in enumerate(project["findings"], start=1):
        report_document.add_heading(
            f"{index}. [{finding['severity']}风险] {finding['title']}（{finding.get('review_status', '待复核')}）",
            level=2,
        )
        report_document.add_paragraph(f"类别：{finding['category']}。{finding['message']}")
        if finding.get("source") or finding.get("location"):
            report_document.add_paragraph(f"原文位置：{finding.get('source', '')} {finding.get('location', '')}")
        if finding.get("context"):
            report_document.add_paragraph(f"原文摘录：{finding['context']}")
        if finding.get("suggestion"):
            report_document.add_paragraph(f"建议：{finding['suggestion']}")
        for citation in finding.get("citations", []):
            report_document.add_paragraph(
                f"制度依据：{citation['title']}；发布方：{citation['issuer']}；来源：{citation['source']}；"
                f"适用范围：{citation['applicability']}；摘录：{citation['excerpt']}"
            )
    report_document.add_heading("三、审查范围与限制", level=1)
    report_document.add_paragraph(
        "内容提取和风险提示基于规则扫描及已导入制度材料。未配置制度来源或 AI 模型的检查项不会伪造法规条款；"
        "OCR、字段抽取与差异比对可能漏报或误报，最终结论须由采购、业务和法务人员结合原件确认。"
    )
    output = BytesIO()
    report_document.save(output)
    output.seek(0)
    report_base = project["name"]
    if Path(report_base).suffix.lower() in ALLOWED_EXTENSIONS:
        report_base = Path(report_base).stem
    filename = quote(f"{report_base}-审查报告.docx")
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{filename}"},
    )


@app.get("/api/knowledge")
def list_knowledge() -> dict[str, Any]:
    with _connect_database() as connection:
        rows = connection.execute(
            "SELECT id, title, category, issuer, source, effective_date, applicability, imported_at, length(text) AS char_count "
            "FROM knowledge ORDER BY imported_at DESC"
        ).fetchall()
    return {"items": [dict(row) for row in rows]}


async def add_knowledge(
    files: list[UploadFile] = File(...),
    category: str = Form("采购制度"),
    issuer: str = Form("待填写"),
    effective_date: str = Form("待核验"),
    applicability: str = Form("请填写适用范围"),
) -> dict[str, Any]:
    imported: list[dict[str, Any]] = []
    for upload in files:
        filename = Path(upload.filename or "未命名制度").name
        raw = await upload.read(MAX_FILE_BYTES + 1)
        await upload.close()
        extracted = await run_in_threadpool(extract_document, filename, raw)
        if extracted.get("error"):
            imported.append({"filename": filename, "error": extracted["error"]})
            continue
        text = extracted.get("text", "").strip()
        if not text:
            imported.append({"filename": filename, "error": "未提取到文本，未加入知识库。"})
            continue
        item_id = uuid.uuid4().hex
        record = {
            "id": item_id,
            "title": Path(filename).stem,
            "category": category.strip() or "采购制度",
            "issuer": issuer.strip() or "待填写",
            "source": filename,
            "effective_date": effective_date.strip() or "待核验",
            "applicability": applicability.strip() or "请填写适用范围",
            "text": text[:MAX_TEXT_CHARS],
            "imported_at": _now(),
        }
        with _connect_database() as connection:
            connection.execute(
                "INSERT INTO knowledge (id,title,category,issuer,source,effective_date,applicability,text,imported_at) "
                "VALUES (:id,:title,:category,:issuer,:source,:effective_date,:applicability,:text,:imported_at)",
                record,
            )
        imported.append({key: value for key, value in record.items() if key != "text"} | {"char_count": len(record["text"])})
    return {"items": imported}


def delete_knowledge(item_id: str) -> dict[str, bool]:
    with _connect_database() as connection:
        cursor = connection.execute("DELETE FROM knowledge WHERE id = ?", (item_id,))
    if not cursor.rowcount:
        raise HTTPException(status_code=404, detail="未找到该制度记录。")
    return {"deleted": True}


@app.get("/api/projects")
def list_projects() -> dict[str, Any]:
    with _connect_database() as connection:
        rows = connection.execute(
            "SELECT id,name,created_at,budget,reviewed,comparison_id,findings_json,documents_json "
            "FROM projects ORDER BY created_at DESC"
        ).fetchall()
    items = []
    for row in rows:
        findings = _read_json(row["findings_json"])
        documents = _read_json(row["documents_json"])
        items.append(
            {
                "id": row["id"],
                "name": row["name"],
                "created_at": row["created_at"],
                "budget": row["budget"],
                "reviewed": bool(row["reviewed"]),
                "comparison_id": row["comparison_id"],
                "document_count": len(documents),
                "finding_count": len(findings),
                "high_count": sum(item["severity"] == "高" for item in findings),
            }
        )
    return {"items": items}


async def review_documents(
    files: list[UploadFile] = File(...),
    project_name: str = Form(""),
    budget: str = Form(""),
    comparison_id: str = Form(""),
) -> dict[str, Any]:
    try:
        budget_override = float(budget.replace(",", "")) if budget.strip() else None
        if budget_override is not None and budget_override < 0:
            raise ValueError
    except ValueError as error:
        raise HTTPException(status_code=422, detail="项目预算必须是非负数字，单位为元。") from error

    documents: list[dict[str, Any]] = []
    for upload in files:
        filename = Path(upload.filename or "未命名文件").name
        raw = await upload.read(MAX_FILE_BYTES + 1)
        await upload.close()
        documents.append(await run_in_threadpool(extract_document, filename, raw))
    if not documents:
        raise HTTPException(status_code=422, detail="请至少上传一个支持的采购文件。")

    fields = _extract_key_fields(documents, budget_override)
    previous = None
    if comparison_id:
        with _connect_database() as connection:
            row = connection.execute("SELECT id,name,fields_json FROM projects WHERE id = ?", (comparison_id,)).fetchone()
        if row:
            previous = {"id": row["id"], "name": row["name"], "fields": _read_json(row["fields_json"])}
    with _connect_database() as connection:
        knowledge_rows = connection.execute("SELECT * FROM knowledge ORDER BY imported_at DESC").fetchall()
    knowledge = [dict(row) for row in knowledge_rows]
    findings = _review_documents(documents, fields, knowledge, budget_override, previous)
    comparison_name = previous["name"] if previous else ""
    stored_documents = [
        {key: document.get(key) for key in ("filename", "format", "char_count", "source_count", "ocr_pages", "error")}
        for document in documents
    ]
    project_id = uuid.uuid4().hex
    resolved_name = project_name.strip() or Path(documents[0]["filename"]).stem
    budget_value = budget_override
    if budget_value is None and fields["budget"]["values"]:
        budget_value = fields["budget"]["values"][0]["value"]
    with _connect_database() as connection:
        connection.execute(
            "INSERT INTO projects (id,name,created_at,budget,fields_json,documents_json,findings_json,comparison_id,reviewed) "
            "VALUES (?,?,?,?,?,?,?,?,0)",
            (
                project_id, resolved_name, _now(), budget_value,
                json.dumps(fields, ensure_ascii=False), json.dumps(stored_documents, ensure_ascii=False),
                json.dumps(findings, ensure_ascii=False), comparison_id,
            ),
        )
    for document in documents:
        document.pop("blocks", None)
        document.pop("text", None)
    return {
        "id": project_id,
        "name": resolved_name,
        "created_at": _now(),
        "budget": budget_value,
        "comparison_name": comparison_name,
        "documents": documents,
        "fields": fields,
        "findings": findings,
        "reviewed": False,
        "ai_status": "仅规则审查；Ollama 未发现已安装模型。" if not _ollama_model() else "已检测到本地模型；当前审查以可追溯规则结果为准。",
    }


def _ollama_model() -> str | None:
    try:
        import urllib.request

        with urllib.request.urlopen("http://127.0.0.1:11434/api/tags", timeout=2) as response:
            payload = json.loads(response.read())
        models = payload.get("models", [])
        return models[0].get("name") if models else None
    except Exception:
        return None


@app.patch("/api/projects/{project_id}")
async def update_project(project_id: str, request: Request) -> dict[str, Any]:
    payload = await request.json()
    findings = payload.get("findings")
    if not isinstance(findings, list):
        raise HTTPException(status_code=422, detail="findings 必须是问题列表。")
    with _connect_database() as connection:
        cursor = connection.execute(
            "UPDATE projects SET findings_json = ?, reviewed = ? WHERE id = ?",
            (json.dumps(findings, ensure_ascii=False), int(bool(payload.get("reviewed"))), project_id),
        )
    if not cursor.rowcount:
        raise HTTPException(status_code=404, detail="未找到该审查项目。")
    return {"id": project_id, "saved": True}


@app.get("/api/projects/{project_id}/report.docx")
def export_report(project_id: str) -> StreamingResponse:
    with _connect_database() as connection:
        row = connection.execute("SELECT * FROM projects WHERE id = ?", (project_id,)).fetchone()
    if not row:
        raise HTTPException(status_code=404, detail="未找到该审查项目。")
    fields = _read_json(row["fields_json"])
    documents = _read_json(row["documents_json"])
    findings = _read_json(row["findings_json"])
    report = Document()
    report.add_heading("采购文档审查报告", level=0)
    report.add_paragraph(f"项目名称：{row['name']}")
    report.add_paragraph(f"审查时间：{row['created_at']}")
    report.add_paragraph(f"项目预算：{row['budget']:,.2f} 元" if row["budget"] is not None else "项目预算：未识别/未填写")
    report.add_paragraph(f"对比项目：{row['comparison_id'] or '未选择'}")
    report.add_heading("一、审查范围", level=1)
    for item in documents:
        report.add_paragraph(f"{item['filename']}（{item.get('format') or '格式未知'}，提取 {item.get('char_count', 0)} 字）", style="List Bullet")
    report.add_heading("二、关键信息摘要", level=1)
    for key, value in fields.items():
        if key == "document_types":
            continue
        label = value.get("label", key)
        values = value.get("values", [])
        if key == "budget" and value.get("manual_value") is not None:
            report.add_paragraph(f"{label}：人工录入 {value['manual_value']:,.2f} 元")
        elif values:
            report.add_paragraph(f"{label}：" + "；".join(str(item.get("display", item["value"]))[:260] for item in values[:4]))
        else:
            report.add_paragraph(f"{label}：未识别")
    report.add_heading("三、问题清单", level=1)
    if not findings:
        report.add_paragraph("规则扫描未生成问题提示；这不代表项目完全合规。")
    for index, item in enumerate(findings, start=1):
        report.add_heading(f"{index}. [{item['severity']}风险] {item['title']}", level=2)
        report.add_paragraph(f"类别：{item['category']}　复核状态：{item.get('review_status', '待复核')}")
        report.add_paragraph(item["message"])
        if item.get("source"):
            report.add_paragraph(f"原文来源：{item['source']}　位置：{item.get('location', '')}")
        if item.get("context"):
            report.add_paragraph(f"原文摘录：{item['context']}")
        if item.get("suggestion"):
            report.add_paragraph(f"修改/核验建议：{item['suggestion']}")
        for citation in item.get("citations", []):
            report.add_paragraph(
                f"参考依据（需人工核验）：{citation['title']}；{citation['issuer']}；"
                f"来源：{citation['source']}；适用范围：{citation['applicability']}；摘录：{citation['excerpt']}"
            )
    report.add_heading("四、审查说明", level=1)
    report.add_paragraph(
        "本报告由本机文档提取与规则检查生成，属于辅助审查提示，不构成法律意见或合规结论。"
        "引用依据仅来自本机知识库中的用户导入资料；适用性、效力状态和最终结论须由人工复核。"
    )
    buffer = BytesIO()
    report.save(buffer)
    buffer.seek(0)
    filename = f"{row['name']}-采购文档审查报告.docx"
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(filename)}"},
    )


if __name__ == "__main__":
    import uvicorn

    uvicorn.run(app, host="127.0.0.1", port=8765)